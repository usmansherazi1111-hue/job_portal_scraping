"""CSV, Excel and JSON exports. Excel has Jobs, Companies (hiring signal per employer) and Sources (counts + terms)."""
from __future__ import annotations
import csv, io, json
from collections import defaultdict
from .model import FIELDS, Job
from .sources import SOURCES

COLS = FIELDS + ["also_on"]


def rows(jobs: list[Job]) -> list[dict]:
    out = []
    for j in jobs:
        r = j.row(); r["also_on"] = ", ".join(j.extra.get("also_on", [])); out.append(r)
    return out


def to_csv(jobs) -> bytes:
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=COLS, extrasaction="ignore"); w.writeheader(); w.writerows(rows(jobs))
    return ("﻿" + buf.getvalue()).encode("utf-8")   # BOM so Excel opens accents correctly


def to_json(jobs, meta: dict) -> bytes:
    return json.dumps({"meta": meta, "jobs": rows(jobs)}, ensure_ascii=False, indent=1).encode("utf-8")


def companies(jobs) -> list[dict]:
    by = defaultdict(list)
    for j in jobs: by[j.company or "(unknown)"].append(j)
    out = []
    for name, js in by.items():
        depts = defaultdict(int)
        for j in js:
            for d in (j.department or "").split(","):
                if d.strip(): depts[d.strip()] += 1
        out.append({"company": name, "open_jobs": len(js), "newest_posting": max((j.posted for j in js), default=""),
                    "top_departments": ", ".join(d for d, _ in sorted(depts.items(), key=lambda x: -x[1])[:4]),
                    "locations": "; ".join(sorted({j.location for j in js if j.location})[:5]),
                    "sources": ", ".join(sorted({j.source for j in js}))})
    return sorted(out, key=lambda r: -r["open_jobs"])


def to_xlsx(jobs, status: list[dict], meta: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook(); head = Font(bold=True, color="FFFFFF"); fill = PatternFill("solid", fgColor="1F3A2E")

    def sheet(ws, header, data, widths):
        ws.append(header)
        for c in ws[1]: c.font, c.fill = head, fill
        for r in data: ws.append([("" if v is None else str(v)[:32000]) for v in r])
        for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions

    ws = wb.active; ws.title = "Jobs"
    R = rows(jobs)
    sheet(ws, COLS, [[r[c] for c in COLS] for r in R], [38, 26, 26, 22, 18, 16, 8, 20, 12, 12, 60, 40, 14, 18, 18])
    for i, r in enumerate(R, 2):
        if r["url"]: ws.cell(i, COLS.index("url") + 1).hyperlink = r["url"]
        ws.cell(i, COLS.index("description") + 1).alignment = Alignment(wrap_text=False)
    C = companies(jobs)
    sheet(wb.create_sheet("Companies"), list(C[0].keys()) if C else ["company"], [list(c.values()) for c in C], [34, 10, 14, 40, 44, 26])
    S = [[s.get("name"), s.get("state"), s.get("found", ""), s.get("kept", ""), s.get("ms", ""), s.get("error", ""),
          (SOURCES.get(s["id"].split(":")[0]).note if SOURCES.get(s["id"].split(":")[0]) else "")] for s in status]
    sheet(wb.create_sheet("Sources"), ["source", "state", "found", "kept", "ms", "message", "terms / attribution"], S, [36, 10, 8, 8, 8, 40, 70])
    ws = wb.create_sheet("About")
    for k, v in meta.items(): ws.append([k, str(v)])
    ws.column_dimensions["A"].width = 18; ws.column_dimensions["B"].width = 90
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()
